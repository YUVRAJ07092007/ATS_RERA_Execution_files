"""Add 'Page X of Y' at top-right of every page in MOM_dated_18062026_DTCP folder."""
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

MOM_DIR = Path(__file__).resolve().parent.parent / "MOM_dated_18062026_DTCP"
HEADER_TEXT = "Page "
HEADER_FONT_SIZE = 9


def _append_field(run, field_code: str) -> None:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def add_docx_page_numbers(path: Path) -> None:
    doc = Document(str(path))
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)

        run = paragraph.add_run(HEADER_TEXT)
        run.font.size = Pt(HEADER_FONT_SIZE)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

        page_run = paragraph.add_run()
        page_run.font.size = Pt(HEADER_FONT_SIZE)
        page_run.font.name = "Times New Roman"
        _append_field(page_run, "PAGE")

        of_run = paragraph.add_run(" of ")
        of_run.font.size = Pt(HEADER_FONT_SIZE)
        of_run.font.name = "Times New Roman"

        total_run = paragraph.add_run()
        total_run.font.size = Pt(HEADER_FONT_SIZE)
        total_run.font.name = "Times New Roman"
        _append_field(total_run, "NUMPAGES")

    doc.save(str(path))


def add_pdf_page_numbers(path: Path) -> None:
    pdf = fitz.open(path)
    total = pdf.page_count
    for index, page in enumerate(pdf):
        rect = page.rect
        label = f"Page {index + 1} of {total}"
        box = fitz.Rect(rect.width - 130, 18, rect.width - 36, 36)
        page.insert_textbox(
            box,
            label,
            fontsize=HEADER_FONT_SIZE,
            fontname="helv",
            align=fitz.TEXT_ALIGN_RIGHT,
        )
    pdf.save(path, incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
    pdf.close()


def main() -> None:
    docx_files = sorted(MOM_DIR.rglob("*.docx"))
    pdf_files = sorted(MOM_DIR.rglob("*.pdf"))

    print(f"Folder: {MOM_DIR}\n")

    for path in docx_files:
        add_docx_page_numbers(path)
        print(f"DOCX: {path.relative_to(MOM_DIR)}")

    for path in pdf_files:
        add_pdf_page_numbers(path)
        print(f"PDF:  {path.relative_to(MOM_DIR)}")

    print(f"\nDone — {len(docx_files)} DOCX, {len(pdf_files)} PDF files updated.")
    print("Note: .md source files are not paginated; use their .docx outputs for filing.")


if __name__ == "__main__":
    main()
