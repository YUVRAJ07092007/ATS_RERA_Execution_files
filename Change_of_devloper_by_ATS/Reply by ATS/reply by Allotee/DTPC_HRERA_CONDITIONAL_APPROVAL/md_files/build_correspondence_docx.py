"""Build DOCX from correspondence markdown (PNB letter, ATS reply, etc.)."""
import re
import sys
import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = "Arial"
BODY_SIZE = 12
PAGE_NUM_FONT_SIZE = 10
SKIP_PREFIXES = ("**Filing note", "**Purpose:", "# Draft for counsel")


def sf(run, *, bold=False, italic=False, size=BODY_SIZE):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_runs(paragraph, text: str, *, bold=False, italic=False, size=BODY_SIZE):
    parts = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_runs(paragraph, part[2:-2], bold=True, italic=italic, size=size)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            sf(run, bold=bold, italic=italic, size=size)
        else:
            run = paragraph.add_run(part)
            sf(run, bold=bold, italic=italic, size=size)


def add_para(doc, text: str, *, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, size=BODY_SIZE):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    add_runs(p, text.strip(), bold=bold, size=size)
    return p


def is_table_row(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|") and "|" in s[1:-1]


def is_table_sep(line: str) -> bool:
    s = line.strip()
    return bool(re.match(r"^\|?[\s\-:|]+\|?$", s))


def _append_field(run, field_code: str) -> None:
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def set_update_fields_on_open(doc: Document) -> None:
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _add_page_number_block(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run("Page ")
    run.font.size = Pt(PAGE_NUM_FONT_SIZE)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    page_run = paragraph.add_run()
    page_run.font.size = Pt(PAGE_NUM_FONT_SIZE)
    page_run.font.name = FONT
    _append_field(page_run, "PAGE")

    of_run = paragraph.add_run(" of ")
    of_run.font.size = Pt(PAGE_NUM_FONT_SIZE)
    of_run.font.name = FONT

    total_run = paragraph.add_run()
    total_run.font.size = Pt(PAGE_NUM_FONT_SIZE)
    total_run.font.name = FONT
    _append_field(total_run, "NUMPAGES")


def add_page_numbers(doc: Document) -> None:
    """Top-right header and bottom-right footer: 'Page X of Y'."""
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        _add_page_number_block(header.paragraphs[0] if header.paragraphs else header.add_paragraph())

        footer = section.footer
        footer.is_linked_to_previous = False
        _add_page_number_block(footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph())


def justify_body_paragraphs(doc: Document) -> None:
    """Ensure body paragraphs use justified alignment (letterhead blocks stay left)."""
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def save_docx_with_retry(doc: Document, docx_path: Path, *, attempts: int = 5) -> None:
    last_error: Exception | None = None
    for attempt in range(1, attempts + 1):
        try:
            doc.save(docx_path)
            return
        except OSError as exc:
            last_error = exc
            if attempt < attempts:
                time.sleep(1.5 * attempt)
    if last_error:
        raise last_error


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, cell_text in enumerate(row):
            cell = table.rows[r].cells[c]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, cell_text.strip(), bold=(r == 0), size=10 if r == 0 else BODY_SIZE)


def build_docx_from_md(text: str, docx_path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    lines = text.splitlines()
    i = 0
    in_filing_note = False
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if not line:
            i += 1
            continue

        if line.startswith("# ") and "Draft for counsel" in line:
            i += 1
            continue

        if any(line.startswith(p) for p in SKIP_PREFIXES):
            if line.startswith("**Filing note"):
                in_filing_note = True
            i += 1
            continue

        if in_filing_note:
            i += 1
            continue

        if line.strip().lower() in ("<!-- pagebreak -->", "<!--pagebreak-->"):
            add_page_break(doc)
            i += 1
            continue

        if line.strip() == "---":
            add_para(doc, "", after=4)
            i += 1
            continue

        if is_table_row(line):
            table_rows: list[list[str]] = []
            while i < len(lines) and is_table_row(lines[i]):
                if not is_table_sep(lines[i]):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    table_rows.append(cells)
                i += 1
            add_table(doc, table_rows)
            add_para(doc, "", after=4)
            continue

        if raw.endswith("  "):
            block = []
            while i < len(lines):
                current_raw = lines[i]
                current = current_raw.rstrip()
                if not current or current.strip() == "---" or is_table_row(current):
                    break
                block.append(current)
                i += 1
                if not current_raw.endswith("  "):
                    break
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(8)
            for idx, bl in enumerate(block):
                if idx > 0:
                    p.add_run().add_break()
                add_runs(p, bl.strip())
            continue

        if line.startswith("## "):
            add_para(doc, line[3:].strip(), bold=True, after=8)
            i += 1
            continue

        add_para(doc, line, after=6)
        i += 1

    justify_body_paragraphs(doc)
    add_page_numbers(doc)
    set_update_fields_on_open(doc)
    save_docx_with_retry(doc, docx_path)


def main() -> None:
    md_dir = Path(__file__).resolve().parent
    letters_dir = Path(__file__).resolve().parents[5] / "PNB_ATS" / "email_letters"
    targets = [
        letters_dir / "23062026_Email_to_PNB_PreEMI_Procedure_Unit_3051.md",
        letters_dir / "23062026_Reply_to_ATS_PreEMI_Meeting_Unit_3051.md",
    ]
    if len(sys.argv) > 1:
        targets = [Path(a) if Path(a).is_absolute() else md_dir / a for a in sys.argv[1:]]

    for md_path in targets:
        md_path = Path(md_path)
        if not md_path.exists():
            print(f"Skip (missing): {md_path}")
            continue
        docx_path = md_path.with_suffix(".docx")
        text = md_path.read_text(encoding="utf-8")
        build_docx_from_md(text, docx_path)
        print(f"Saved: {docx_path}")


if __name__ == "__main__":
    main()
