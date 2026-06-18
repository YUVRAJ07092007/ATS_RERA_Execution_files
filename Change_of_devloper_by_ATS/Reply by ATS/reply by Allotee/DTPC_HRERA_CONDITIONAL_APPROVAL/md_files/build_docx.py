"""Build DOCX from master consolidated issues list markdown."""
import re
from pathlib import Path

import markdown
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from htmldocx import HtmlToDocx

MD_PATH = Path(__file__).with_name('00_MASTER_CONSOLIDATED_ISSUES_LIST_COD_AND_EXECUTION_Unit_3051.md')
DOCX_PATH = MD_PATH.parent.parent / '00_MASTER_CONSOLIDATED_ISSUES_LIST_COD_AND_EXECUTION_Unit_3051.docx'

FONT_NAME = 'Times New Roman'
FONT_SIZE = Pt(12)
HEADING_STYLES = {'Heading 1', 'Heading 2', 'Heading 3'}

DOCX_OMIT_PATTERNS = (
    'working folder:',
)

TABLE_TITLES = (
    'Key terms — abbreviations and full forms',
    'Three separate tracks',
    'HARERA order 18.04.2024 — binding technical directions (Unit 3051)',
    'DTCP in-principle COD (10.12.2025) — Conditions 1–12',
    '§1 Entity identification and roles — requirements',
    '§1A Intra-group restructuring — precautionary issues',
    '§2 Asset transfer — requirements',
    '§3 Liability transfer and assumption — requirements',
    '§4 COD terms, agreements and DTCP conditions — requirements',
    '§5 Board resolutions, POA and deponent authority — requirements',
    '§6 Impugned joint affidavit dated 08.06.2026 — requirements',
    '§6A Possession date — affidavit gap and 30.06.2031 projection — requirements',
    '§6B HARERA order compliance matrix',
    '§7 Unit 3051 ledger, payments and accounts — requirements',
    '§7A COD and subvention / pre-EMI — precautionary issues',
    '§8 Allottee list, notice and objection record — requirements',
    '§9 Takeover / handling admission by ATS group — requirements',
    '§10 Core compliance safeguards for Unit 3051 — requirements',
    '§11 Joint liability on DTCP / COD file — requirements',
    '§12 Plan approval and objection process — requirements',
    '§13 Possession, defect liability and coercive practices — requirements',
    '§14 Escrow, project funding and financial disclosure — requirements',
    '§15 Regulatory process, hearings and service — requirements',
    '§16 ATS conduct, admissions and clarifications — requirements',
)

TABLE_CAPTION_RE = re.compile(r'^\*\*Table\s+(\d+)\s*[—–-]', re.IGNORECASE)


def is_table_start(lines: list[str], idx: int) -> bool:
    line = lines[idx].strip()
    if not line.startswith('|'):
        return False
    if idx + 1 >= len(lines):
        return False
    next_line = lines[idx + 1].strip()
    return next_line.startswith('|') and '---' in next_line


def caption_before(lines: list[str], idx: int) -> str | None:
    for j in range(idx - 1, max(-1, idx - 6), -1):
        stripped = lines[j].strip()
        if not stripped:
            continue
        match = TABLE_CAPTION_RE.match(stripped)
        if match:
            return stripped
        return None
    return None


def insert_table_captions(lines: list[str]) -> list[str]:
    result: list[str] = []
    table_num = 0
    i = 0
    while i < len(lines):
        if is_table_start(lines, i) and not caption_before(result, len(result)):
            table_num += 1
            title = TABLE_TITLES[table_num - 1] if table_num <= len(TABLE_TITLES) else 'Requirements'
            if result and result[-1].strip():
                result.append('')
            result.append(f'**Table {table_num} — {title}**')
            result.append('')
        result.append(lines[i])
        i += 1
    return result


def omit_from_docx(line: str) -> bool:
    normalized = line.strip().lower()
    if any(pattern in normalized for pattern in DOCX_OMIT_PATTERNS):
        return True
    return 'place:' in normalized and 'new delhi' in normalized


def apply_run_font(run, *, heading: bool = False) -> None:
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run.font.color.rgb = RGBColor(0, 0, 0)
    if heading:
        run.font.bold = True


def add_page_number(paragraph) -> None:
    """Insert a proper Word PAGE field (updates per page when opened in Word)."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run_begin = paragraph.add_run()
    r = run_begin._r
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    r.append(fld_begin)

    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' PAGE '
    r.append(instr)

    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    r.append(fld_sep)

    run_num = paragraph.add_run('1')
    apply_run_font(run_num)

    run_end = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run_end._r.append(fld_end)

    apply_run_font(run_begin)
    apply_run_font(run_end)


def apply_paragraph_alignment(paragraph, *, heading: bool = False) -> None:
    if heading:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def style_paragraphs(doc: Document) -> None:
    normal = doc.styles['Normal']
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name in HEADING_STYLES | {'List Bullet', 'List Number'}:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT_NAME
            style.font.size = FONT_SIZE
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if style_name in HEADING_STYLES:
                style.font.bold = True

    for para in doc.paragraphs:
        is_heading = para.style.name in HEADING_STYLES
        apply_paragraph_alignment(para, heading=is_heading)
        for run in para.runs:
            apply_run_font(run, heading=is_heading)


def style_tables(doc: Document) -> None:
    for table in doc.tables:
        hdr = table.rows[0].cells[0].text.strip().lower() if table.rows else ''
        headers = [c.text.strip().lower() for c in table.rows[0].cells] if table.rows else []

        if len(table.columns) >= 2:
            if hdr == 'abbrev.' and 'full form' in headers:
                table.columns[0].width = Cm(2.7)
                table.columns[1].width = Cm(4.5)
                table.columns[2].width = Cm(9.8)
            elif hdr == 'condition' and 'technical subject' in headers:
                table.columns[0].width = Cm(2.0)
                table.columns[1].width = Cm(14.0)
            elif hdr in ('track', 'direction', 'harera direction'):
                w = Cm(16.0 / len(table.columns))
                for i in range(len(table.columns)):
                    table.columns[i].width = w
            elif hdr == 'no.' and 'what is required' in headers:
                table.columns[0].width = Cm(1.5)
                table.columns[1].width = Cm(14.5)
            elif hdr == 'priority' and len(table.columns) == 3:
                table.columns[0].width = Cm(1.5)
                table.columns[1].width = Cm(2.5)
                table.columns[2].width = Cm(12.0)
            elif hdr in ('column', 'annexure code'):
                table.columns[0].width = Cm(2.5)
                table.columns[1].width = Cm(13.5)

        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_before = Pt(0)
                    para.paragraph_format.space_after = Pt(0)
                    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    for run in para.runs:
                        apply_run_font(run)


def style_table_captions(doc: Document) -> None:
    for para in doc.paragraphs:
        text = para.text.strip()
        if re.match(r'^Table\s+\d+\s*[—–-]', text, re.IGNORECASE):
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(3)
            for run in para.runs:
                apply_run_font(run, heading=True)


def build() -> None:
    text = MD_PATH.read_text(encoding='utf-8')
    lines = insert_table_captions(text.splitlines())
    processed = []
    for line in lines:
        if omit_from_docx(line):
            continue
        if line.strip().startswith('|') and (not processed or not processed[-1].strip().startswith('|')):
            if processed and processed[-1].strip() != '':
                processed.append('')
        processed.append(line)

    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.header_distance = Cm(0.8)
        section.footer_distance = Cm(0.8)

        hp = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
        hp.clear()
        add_page_number(hp)

    md_html = markdown.markdown('\n'.join(processed), extensions=['tables'])
    HtmlToDocx().add_html_to_document(md_html, doc)

    style_paragraphs(doc)
    style_table_captions(doc)
    style_tables(doc)
    doc.save(DOCX_PATH)
    print('Saved:', DOCX_PATH)


if __name__ == '__main__':
    import sys

    if len(sys.argv) > 1 and sys.argv[1] == '--update-md':
        text = MD_PATH.read_text(encoding='utf-8')
        updated = '\n'.join(insert_table_captions(text.splitlines())) + '\n'
        MD_PATH.write_text(updated, encoding='utf-8')
        print('Updated:', MD_PATH)
    else:
        build()
