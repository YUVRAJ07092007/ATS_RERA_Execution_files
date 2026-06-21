"""Regenerate affidavit DOCX from markdown after every .md update."""
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

MD_PATH = Path(__file__).with_name("20062026_DRAFT_Affidavit_cum_Undertaking_Unit_3051.md")
DOCX_PATH = MD_PATH.parent.parent / "20062026_DRAFT_Affidavit_cum_Undertaking_Unit_3051.docx"

FONT = "Times New Roman"
BODY_SIZE = 11
TIMESTAMP_RE = re.compile(r"^\*\*Last updated:\*\*\s*.+\n?", re.MULTILINE)


def now_stamp() -> str:
    return datetime.now().strftime("%d.%m.%Y %H:%M:%S")


def stamp_line(stamp: str | None = None) -> str:
    return f"**Last updated:** {stamp or now_stamp()}\n"


def update_md_timestamp(text: str, stamp: str | None = None) -> tuple[str, str]:
    stamp = stamp or now_stamp()
    line = stamp_line(stamp)
    if TIMESTAMP_RE.search(text):
        text = TIMESTAMP_RE.sub(line, text, count=1)
    else:
        text = text.replace("\n", f"\n{line}", 1)
    return text, stamp


def strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"\*(.+?)\*", r"\1", text)
    return text.strip()


def sf(run, *, bold=False, size=BODY_SIZE):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold


def add_para(doc, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, size=BODY_SIZE):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    sf(p.add_run(strip_md(text)), bold=bold, size=size)


def parse_md(text: str) -> dict:
    stamp_match = re.search(r"\*\*Last updated:\*\*\s*(.+)", text)
    timestamp = stamp_match.group(1).strip() if stamp_match else now_stamp()

    date_match = re.search(r"\*\*Date:\*\*\s*(.+)", text)
    date = date_match.group(1).strip() if date_match else ""

    start = text.find("## AFFIDAVIT-CUM-UNDERTAKING")
    end = text.find("**Filing note:**")
    body = text[start:end] if start != -1 else text

    before_match = re.search(
        r"\*\*Before:\*\*\s*(.+?)\n\n---",
        body,
        re.DOTALL,
    )
    before = strip_md(before_match.group(1).replace("\n", " ")) if before_match else ""

    intro_match = re.search(
        r"We, the deponents hereinbelow.+?\n\n---",
        body,
        re.DOTALL,
    )
    intro = strip_md(intro_match.group(0).replace("\n---", "")) if intro_match else ""

    def section_block(header: str) -> tuple[str, str, list[str]]:
        pattern = rf"## {re.escape(header)}\s*\n\*(.+?)\*\n\n"
        m = re.search(pattern, body, re.DOTALL)
        subtitle = strip_md(m.group(1)) if m else ""
        start_idx = m.end() if m else body.find(f"## {header}")
        next_headers = ["## PART II", "## PART I", "### Verification", "---"]
        end_idx = len(body)
        for nh in next_headers:
            if nh == f"## {header}":
                continue
            pos = body.find(nh, start_idx)
            if pos != -1:
                end_idx = min(end_idx, pos)
        chunk = body[start_idx:end_idx]
        paras = []
        for line in chunk.splitlines():
            line = line.strip()
            if not line or line == "---":
                continue
            if line.startswith("*") and line.endswith("*"):
                continue
            paras.append(line)
        return header, subtitle, paras

    _, part1_sub, part1 = section_block("PART I — Text of joint affidavit dated 08.06.2026 (verbatim)")
    _, part2_sub, part2 = section_block("PART II — Unit No. 3051 — specific undertakings")

    verif_match = re.search(r"### Verification\n\n(.+?)\n\n---", body, re.DOTALL)
    verification = strip_md(verif_match.group(1)) if verif_match else ""

    sig_blocks = []
    sig_section = re.search(
        r"### Verification\n\n.+?\n\n---\n\n(.+?)\n\n---\n\n\*\*Before me",
        body,
        re.DOTALL,
    )
    if sig_section:
        for block in re.split(r"\n\n(?=\*\*For )", sig_section.group(1).strip()):
            lines = [strip_md(x) for x in block.splitlines() if x.strip()]
            if lines and lines[0].startswith("For "):
                sig_blocks.append((lines[0][4:], lines[1:]))

    notary_match = re.search(
        r"\*\*Before me, Notary Public / Oath Commissioner\*\*\n\n(.+?)\n\n---",
        body,
        re.DOTALL,
    )
    notary = strip_md(notary_match.group(1)) if notary_match else ""

    enclosures = []
    enc_section = re.search(r"### Enclosures\n\n(.+?)\n\n---", body, re.DOTALL)
    if enc_section:
        for line in enc_section.group(1).splitlines():
            line = line.strip()
            if line:
                enclosures.append(strip_md(line))

    return {
        "timestamp": timestamp,
        "date": date,
        "before": before,
        "intro": intro,
        "part1_sub": part1_sub,
        "part1": part1,
        "part2_sub": part2_sub,
        "part2": part2,
        "verification": verification,
        "sig_blocks": sig_blocks,
        "notary": notary,
        "enclosures": enclosures,
    }


def build_docx(data: dict) -> None:
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(4.5)
        s.bottom_margin = Cm(2.5)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2.5)

    add_para(doc, f"Last updated: {data['timestamp']}", align=WD_ALIGN_PARAGRAPH.LEFT, after=6, size=10)
    add_para(doc, f"Date: {data['date']}", align=WD_ALIGN_PARAGRAPH.LEFT, after=12)
    add_para(doc, "AFFIDAVIT-CUM-UNDERTAKING", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, size=12)
    add_para(
        doc,
        "Specific to Unit No. 3051 — ATS Grandstand Phase-I, Sector 99A, Gurugram",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    add_para(
        doc,
        "(Part I: generic 08.06.2026 only  |  Part II: Unit 3051 specific)",
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=12,
        size=10,
    )
    add_para(doc, f"Before: {data['before']}", align=WD_ALIGN_PARAGRAPH.LEFT, after=12)
    if data["intro"]:
        add_para(doc, data["intro"], after=10)

    add_para(doc, "PART I — Text of joint affidavit dated 08.06.2026 (verbatim)", bold=True, after=4)
    add_para(doc, f"({data['part1_sub']})", after=8, size=10)
    for p in data["part1"]:
        add_para(doc, p, after=6)

    add_para(doc, "PART II — Unit No. 3051 — specific undertakings", bold=True, after=10)
    add_para(doc, f"({data['part2_sub']})", after=8, size=10)
    for p in data["part2"]:
        add_para(doc, p, after=6)

    add_para(doc, "Verification", bold=True, after=8)
    add_para(doc, data["verification"], after=12)

    for company, lines in data["sig_blocks"]:
        add_para(doc, f"For {company}", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
        for line in lines:
            add_para(doc, line, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
        add_para(doc, "", after=8)

    add_para(doc, "Before me, Notary Public / Oath Commissioner", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    add_para(doc, data["notary"], align=WD_ALIGN_PARAGRAPH.LEFT, after=8)

    add_para(doc, "Enclosures:", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=4)
    for enc in data["enclosures"]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        sf(p.add_run(enc))

    doc.save(DOCX_PATH)


def main() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    text, stamp = update_md_timestamp(text)
    MD_PATH.write_text(text, encoding="utf-8")
    data = parse_md(text)
    build_docx(data)
    print(f"Timestamp: {stamp}")
    print(f"Updated MD:  {MD_PATH}")
    print(f"Refreshed:   {DOCX_PATH}")


if __name__ == "__main__":
    main()
