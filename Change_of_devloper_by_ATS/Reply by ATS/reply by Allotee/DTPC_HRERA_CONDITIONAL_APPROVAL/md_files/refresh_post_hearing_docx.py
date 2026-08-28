"""Regenerate Post-Hearing Representation DOCX from markdown — faithful conversion."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

MOM_DIR = Path(__file__).resolve().parent.parent / "MOM_dated_18062026_DTCP"
MD_PATH = MOM_DIR / "19062026_Post_Hearing_Representation_Unit_3051.md"
DOCX_PATH = MOM_DIR / "19062026_Post_Hearing_Representation_Unit_3051.docx"

FONT = "Times New Roman"
BODY_SIZE = 12

CENTER_HEADERS = (
    "Post-Hearing Representation",
)


def sf(run, *, bold=False, italic=False, size=BODY_SIZE):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_runs_from_md(paragraph, text: str, *, bold=False, italic=False, size=BODY_SIZE):
    parts = re.split(r"(\*\*.+?\*\*|`[^`]+`|\*.+?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_runs_from_md(paragraph, part[2:-2], bold=True, italic=italic, size=size)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            sf(run, bold=bold, italic=italic, size=size)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            add_runs_from_md(paragraph, part[1:-1], bold=bold, italic=True, size=size)
        else:
            run = paragraph.add_run(part)
            sf(run, bold=bold, italic=italic, size=size)


def add_para(
    doc,
    text: str,
    *,
    bold=False,
    italic=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    after=6,
    size=BODY_SIZE,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    add_runs_from_md(p, text.strip(), bold=bold, italic=italic, size=size)
    return p


def add_para_with_line_breaks(
    doc,
    block_lines: list[str],
    *,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    after=8,
    size=BODY_SIZE,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    for idx, line_text in enumerate(block_lines):
        if idx > 0:
            p.add_run().add_break()
        add_runs_from_md(p, line_text.strip(), size=size)
    return p


def is_center_header(text: str) -> bool:
    plain = re.sub(r"\*+", "", text).strip()
    return any(plain.startswith(prefix) for prefix in CENTER_HEADERS)


def heading_level(line: str) -> int | None:
    if line.startswith("### "):
        return 3
    if line.startswith("## "):
        return 2
    if line.startswith("# "):
        return 1
    return None


def build_docx_from_md(text: str, docx_path: Path) -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    lines = text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if not line:
            i += 1
            continue

        if line.strip() == "---":
            add_para(doc, "", after=4)
            i += 1
            continue

        level = heading_level(line)
        if level in (1, 2, 3):
            title = line.lstrip("#").strip()
            align = (
                WD_ALIGN_PARAGRAPH.CENTER
                if is_center_header(title)
                else WD_ALIGN_PARAGRAPH.JUSTIFY
            )
            add_para(
                doc,
                title,
                bold=True,
                align=align,
                after=8,
                size=12 if is_center_header(title) else BODY_SIZE,
            )
            i += 1
            continue

        if raw.endswith("  "):
            block_lines = []
            while i < len(lines):
                current_raw = lines[i]
                current = current_raw.rstrip()
                if not current or current.strip() == "---":
                    break
                if heading_level(current):
                    break
                block_lines.append(current)
                i += 1
                if not current_raw.endswith("  "):
                    break
            add_para_with_line_breaks(doc, block_lines)
            continue

        add_para(doc, line, after=6)
        i += 1

    doc.save(docx_path)


def main() -> None:
    from add_page_numbers_mom_folder import add_docx_page_numbers

    md_path = MD_PATH
    docx_path = DOCX_PATH
    if len(sys.argv) > 1:
        md_path = Path(sys.argv[1])
        docx_path = md_path.with_suffix(".docx")

    text = md_path.read_text(encoding="utf-8")
    build_docx_from_md(text, docx_path)
    add_docx_page_numbers(docx_path)
    print(f"Source MD:  {md_path}")
    print(f"Refreshed:  {docx_path}")


if __name__ == "__main__":
    main()
